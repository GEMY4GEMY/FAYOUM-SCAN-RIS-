from flask import Flask,request,jsonify,send_from_directory,send_file
import sqlite3,os,sys,secrets,hashlib,json,shutil
from datetime import datetime
from io import BytesIO
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side

BUNDLE=getattr(sys,"_MEIPASS",os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RUNTIME=os.path.dirname(sys.executable) if getattr(sys,"frozen",False) else os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB=os.path.join(RUNTIME,"data","fayoum_scan.db"); BACKUPS=os.path.join(RUNTIME,"backups")
SCHEMA=os.path.join(BUNDLE,"server","schema.sql"); WEB=os.path.join(BUNDLE,"web")
app=Flask(__name__,static_folder=WEB); TOKENS={}

def con():
 c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; c.execute("PRAGMA foreign_keys=ON"); return c
def hp(p,s=None):
 s=s or secrets.token_hex(16); return s+"$"+hashlib.pbkdf2_hmac("sha256",p.encode(),s.encode(),200000).hex()
def vp(p,h):
 try:s,x=h.split("$",1);return secrets.compare_digest(hp(p,s).split("$",1)[1],x)
 except:return False
def user():
 return TOKENS.get(request.headers.get("Authorization","").replace("Bearer ",""))
def admin(): return user() if user() and user()["role"]=="admin" else None
def init():
 os.makedirs(os.path.dirname(DB),exist_ok=True);os.makedirs(BACKUPS,exist_ok=True)
 with con() as c:
  c.executescript(open(SCHEMA,encoding="utf-8").read())
  if not c.execute("select count(*) from users").fetchone()[0]:
   c.execute("insert into users(username,password_hash,role) values(?,?,?)",("admin",hp("admin123"),"admin"))

@app.get("/api/health")
def health():
 with con() as c:return jsonify(ok=True,server="FAYOUM SCAN RIS Phase 5",patients=c.execute("select count(*) from patients").fetchone()[0])
@app.post("/api/login")
def login():
 d=request.json or {}
 with con() as c:
  r=c.execute("select * from users where username=? and active=1",(d.get("username"),)).fetchone()
  if not r or not vp(d.get("password",""),r["password_hash"]):return jsonify(error="invalid credentials"),401
  t=secrets.token_urlsafe(32);TOKENS[t]={"id":r["id"],"username":r["username"],"role":r["role"]}
  c.execute("update users set last_login=CURRENT_TIMESTAMP where id=?",(r["id"],));return jsonify(token=t,user=TOKENS[t])
@app.post("/api/logout")
def logout(): TOKENS.pop(request.headers.get("Authorization","").replace("Bearer ",""),None);return jsonify(ok=True)
@app.get("/api/users")
def users():
 if not admin():return jsonify(error="admin required"),403
 with con() as c:return jsonify([dict(x) for x in c.execute("select id,username,role,active,last_login,created_at from users order by username")])
@app.post("/api/users")
def adduser():
 if not admin():return jsonify(error="admin required"),403
 d=request.json or {}
 with con() as c:
  try:r=c.execute("insert into users(username,password_hash,role) values(?,?,?)",(d["username"],hp(d["password"]),d.get("role","user")));return jsonify(id=r.lastrowid),201
  except sqlite3.IntegrityError:return jsonify(error="username exists"),409
@app.get("/api/payment-types")
def pays():
 with con() as c:return jsonify([dict(x) for x in c.execute("select * from payment_types where active=1 order by name")])
@app.get("/api/exams")
def exams():
 u=request.args.get("unit");q="select * from exams where active=1";a=[]
 if u:q+=" and unit=?";a=[u]
 with con() as c:return jsonify([dict(x) for x in c.execute(q+" order by unit,name",a)])
@app.post("/api/payment-types")
def addpay():
 if not admin():return jsonify(error="admin required"),403
 d=request.json or {}
 try:
  with con() as c:r=c.execute("insert into payment_types(name,category,notes) values(?,?,?)",(d["name"],d.get("category"),d.get("notes")));return jsonify(id=r.lastrowid),201
 except sqlite3.IntegrityError:return jsonify(error="payment type exists"),409
@app.post("/api/exams")
def addexam():
 if not admin():return jsonify(error="admin required"),403
 d=request.json or {}
 try:
  with con() as c:r=c.execute("insert into exams(unit,name,base_price) values(?,?,?)",(d["unit"],d["name"],float(d.get("base_price",0))));return jsonify(id=r.lastrowid),201
 except sqlite3.IntegrityError:return jsonify(error="exam exists"),409
@app.get("/api/price-plans")
def plans():
 with con() as c:return jsonify([dict(x) for x in c.execute("select p.*,t.name payment_type from price_plans p join payment_types t on t.id=p.payment_type_id order by p.id desc")])
@app.post("/api/price-plans")
def addplan():
 if not admin():return jsonify(error="admin required"),403
 d=request.json or {}
 with con() as c:r=c.execute("insert into price_plans(payment_type_id,name,version,valid_from,valid_to,active) values(?,?,?,?,?,?)",(d["payment_type_id"],d["name"],int(d.get("version",1)),d["valid_from"],d.get("valid_to"),int(d.get("active",1))));return jsonify(id=r.lastrowid),201
@app.post("/api/price-plans/<int:pid>/items")
def planitem(pid):
 if not admin():return jsonify(error="admin required"),403
 d=request.json or {}
 with con() as c:c.execute("insert into price_plan_items(price_plan_id,exam_id,price,coverage_percentage,active) values(?,?,?,?,1) on conflict(price_plan_id,exam_id) do update set price=excluded.price,coverage_percentage=excluded.coverage_percentage,active=1",(pid,d["exam_id"],float(d["price"]),float(d.get("coverage_percentage",0))))
 return jsonify(ok=True)
@app.get("/api/pricing/resolve")
def resolve():
 pt=request.args.get("payment_type","");ex=request.args.get("exam","");day=request.args.get("date") or datetime.now().date().isoformat()
 with con() as c:
  r=c.execute("select i.id item_id,p.id plan_id,p.name plan_name,p.version,i.price,i.coverage_percentage from price_plan_items i join price_plans p on p.id=i.price_plan_id join payment_types t on t.id=p.payment_type_id join exams e on e.id=i.exam_id where t.name=? and e.name=? and p.active=1 and i.active=1 and p.valid_from<=? and (p.valid_to is null or p.valid_to='' or p.valid_to>=?) order by p.version desc,p.id desc limit 1",(pt,ex,day,day)).fetchone()
  if r:return jsonify(found=True,**dict(r))
  e=c.execute("select base_price from exams where name=? and active=1 order by id desc limit 1",(ex,)).fetchone()
  return jsonify(found=False,price=(e["base_price"] if e else 0),coverage_percentage=0)
@app.get("/api/patients")
def patients():
 q="select * from patients where 1=1";a=[]
 for col,arg in [("case_date","date"),("unit","unit")]:
  v=request.args.get(arg)
  if v and v!="all":q+=f" and {col}=?";a.append(v)
 s=request.args.get("search","")
 if s:q+=" and (name like ? or phone like ? or exam like ? or doctor like ?)";a += [f"%{s}%"]*4
 with con() as c:return jsonify([dict(x) for x in c.execute(q+" order by id desc limit 3000",a)])
@app.post("/api/patients")
def patientadd():
 d=request.json or {}; cols=["name","payment_type","phone","exam","exam_price","coverage_percentage","coverage_amount","additional_fees","discount","total_amount","doctor","notes","username","unit","case_date","timestamp","price_plan_id","price_plan_item_id","price_snapshot"]
 with con() as c:
  r=c.execute(f"insert into patients({','.join(cols)}) values({','.join(['?']*len(cols))})",[d.get(k) for k in cols])
  c.execute("insert into audit_log(username,action,entity_type,entity_id,new_data) values(?,?,?,?,?)",(d.get("username"),"CREATE","patient",r.lastrowid,json.dumps(d,ensure_ascii=False)));return jsonify(id=r.lastrowid),201
@app.put("/api/patients/<int:i>")
def patientedit(i):
 d=request.json or {};expected=int(d.get("row_version",1));u=user();allowed=["name","payment_type","phone","exam","exam_price","coverage_percentage","coverage_amount","additional_fees","discount","total_amount","doctor","notes","unit","case_date"]
 with con() as c:
  old=c.execute("select * from patients where id=?",(i,)).fetchone()
  if not old:return jsonify(error="patient not found"),404
  vals=[d.get(k,old[k]) for k in allowed]+[((u or {}).get("username") or d.get("username")),i,expected]
  cur=c.execute("update patients set "+",".join(k+"=?" for k in allowed)+",username=?,row_version=row_version+1,updated_at=CURRENT_TIMESTAMP where id=? and row_version=?",vals)
  if not cur.rowcount:return jsonify(error="record changed by another user"),409
  c.execute("insert into audit_log(username,action,entity_type,entity_id,old_data,new_data) values(?,?,?,?,?,?)",((u or {}).get("username"),"UPDATE","patient",i,json.dumps(dict(old),ensure_ascii=False),json.dumps(d,ensure_ascii=False)))
 return jsonify(ok=True)
@app.delete("/api/patients/<int:i>")
def patientdelete(i):
 a=admin()
 if not a:return jsonify(error="admin required"),403
 with con() as c:
  old=c.execute("select * from patients where id=?",(i,)).fetchone()
  if not old:return jsonify(error="patient not found"),404
  c.execute("delete from patients where id=?",(i,));c.execute("insert into audit_log(username,action,entity_type,entity_id,old_data) values(?,?,?,?,?)",(a["username"],"DELETE","patient",i,json.dumps(dict(old),ensure_ascii=False)))
 return jsonify(ok=True)
@app.post("/api/migration/rev7")
def migrate():
 if not admin():return jsonify(error="admin required"),403
 d=request.json or {};stats={}
 with con() as c:
  for p in d.get("paymentTypes",[]):
   try:c.execute("insert into payment_types(name,category,notes) values(?,?,?)",(p.get("name"),p.get("category"),p.get("notes")))
   except sqlite3.IntegrityError:pass
  for e in d.get("exams",[]):
   try:c.execute("insert into exams(unit,name,base_price) values(?,?,?)",(e.get("unit"),e.get("name"),float(e.get("price") or 0)))
   except sqlite3.IntegrityError:pass
  inserted=0
  for p in d.get("patients",[]):
   if p.get("id") is not None and c.execute("select 1 from patients where legacy_id=?",(p.get("id"),)).fetchone():continue
   cols=["legacy_id","name","payment_type","phone","exam","exam_price","coverage_percentage","coverage_amount","additional_fees","discount","total_amount","doctor","notes","username","unit","case_date","timestamp"]
   vals=[p.get("id"),p.get("name"),p.get("paymentType"),p.get("phone"),p.get("exam"),p.get("examPrice"),p.get("coveragePercentage"),p.get("coverageAmount"),p.get("additionalFees"),p.get("discount"),p.get("totalAmount"),p.get("doctor"),p.get("notes"),p.get("user"),p.get("unit"),p.get("date"),p.get("timestamp")]
   c.execute("insert into patients("+",".join(cols)+") values("+",".join(["?"]*len(cols))+")",vals);inserted+=1
  stats={"patients_inserted":inserted,"patients_received":len(d.get("patients",[])),"exams_received":len(d.get("exams",[])),"payment_types_received":len(d.get("paymentTypes",[]))}
 return jsonify(ok=True,**stats)
@app.get("/api/appointments")
def appts():
 q="select * from appointments where 1=1";a=[]
 if request.args.get("date"):q+=" and appointment_date=?";a.append(request.args["date"])
 if request.args.get("status") not in (None,"all"):q+=" and status=?";a.append(request.args["status"])
 with con() as c:return jsonify([dict(x) for x in c.execute(q+" order by appointment_date,appointment_time,id desc",a)])
@app.post("/api/appointments")
def apptadd():
 d=request.json or {};u=user();cols=["patient_name","phone","payment_type","unit","exam","appointment_date","appointment_time","period","doctor","notes","status","created_by","updated_by"];v=[d.get(k) for k in cols];v[10]=v[10] or "Booked";v[11]=v[11] or (u["username"] if u else "");v[12]=v[11]
 with con() as c:r=c.execute(f"insert into appointments({','.join(cols)}) values({','.join(['?']*len(cols))})",v);return jsonify(id=r.lastrowid),201
@app.put("/api/appointments/<int:i>")
def apptedit(i):
 d=request.json or {};u=user();expected=int(d.get("row_version",1));allowed=["patient_name","phone","payment_type","unit","exam","appointment_date","appointment_time","period","doctor","notes","status"]
 with con() as c:
  old=c.execute("select * from appointments where id=?",(i,)).fetchone()
  if not old:return jsonify(error="appointment not found"),404
  vals=[d.get(k,old[k]) for k in allowed]+[((u or {}).get("username")),i,expected]
  cur=c.execute("update appointments set "+",".join(k+"=?" for k in allowed)+",updated_by=?,row_version=row_version+1,updated_at=CURRENT_TIMESTAMP where id=? and row_version=?",vals)
  if not cur.rowcount:return jsonify(error="appointment changed by another user"),409
  c.execute("insert into audit_log(username,action,entity_type,entity_id,old_data,new_data) values(?,?,?,?,?,?)",((u or {}).get("username"),"UPDATE","appointment",i,json.dumps(dict(old),ensure_ascii=False),json.dumps(d,ensure_ascii=False)))
 return jsonify(ok=True)
@app.post("/api/appointments/<int:i>/convert")
def apptconvert(i):
 u=user()
 with con() as c:
  a=c.execute("select * from appointments where id=?",(i,)).fetchone()
  if not a:return jsonify(error="appointment not found"),404
  if a["converted_patient_id"]:return jsonify(error="already converted",patient_id=a["converted_patient_id"]),409
  r=c.execute("insert into patients(name,payment_type,phone,exam,exam_price,coverage_percentage,coverage_amount,additional_fees,discount,total_amount,doctor,notes,username,unit,case_date,timestamp) values(?,?,?,?,0,0,0,0,0,0,?,?,?,?,?,?)",(a["patient_name"],a["payment_type"] or "",a["phone"],a["exam"] or "",a["doctor"] or "",a["notes"] or "",(u or {}).get("username"),a["unit"],a["appointment_date"],datetime.now().isoformat()))
  c.execute("update appointments set status='Converted',converted_patient_id=?,updated_by=?,row_version=row_version+1,updated_at=CURRENT_TIMESTAMP where id=?",(r.lastrowid,(u or {}).get("username"),i))
  c.execute("insert into audit_log(username,action,entity_type,entity_id,new_data) values(?,?,?,?,?)",((u or {}).get("username"),"CONVERT","appointment",i,json.dumps({"patient_id":r.lastrowid})))
 return jsonify(ok=True,patient_id=r.lastrowid)
@app.delete("/api/appointments/<int:i>")
def apptdelete(i):
 a=admin()
 if not a:return jsonify(error="admin required"),403
 with con() as c:
  old=c.execute("select * from appointments where id=?",(i,)).fetchone()
  if not old:return jsonify(error="appointment not found"),404
  c.execute("delete from appointments where id=?",(i,));c.execute("insert into audit_log(username,action,entity_type,entity_id,old_data) values(?,?,?,?,?)",(a["username"],"DELETE","appointment",i,json.dumps(dict(old),ensure_ascii=False)))
 return jsonify(ok=True)
@app.get("/api/export/patients.xlsx")
def exportpatients():
 q="select * from patients where 1=1";a=[]
 for col,arg in [("case_date","date"),("unit","unit")]:
  v=request.args.get(arg)
  if v and v!="all":q+=" and "+col+"=?";a.append(v)
 with con() as c:rows=[dict(x) for x in c.execute(q+" order by id",a)]
 wb=Workbook();ws=wb.active;ws.title="Patients"
 headers=["ID","Date","Name","Phone","Payment Type","Unit","Exam","Exam Price","Coverage %","Coverage Amount","Additional Fees","Discount","Total","Doctor","User"]
 keys=["id","case_date","name","phone","payment_type","unit","exam","exam_price","coverage_percentage","coverage_amount","additional_fees","discount","total_amount","doctor","username"]
 ws.append(headers)
 for r in rows:ws.append([r.get(k) for k in keys])
 fill=PatternFill("solid",fgColor="1B2A6B");font=Font(color="FFFFFF",bold=True)
 for cell in ws[1]:cell.fill=fill;cell.font=font;cell.alignment=Alignment(horizontal="center")
 widths=[8,14,28,16,22,14,28,14,12,16,16,12,14,24,18]
 for i,w in enumerate(widths,1):ws.column_dimensions[chr(64+i)].width=w
 ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
 out=BytesIO();wb.save(out);out.seek(0)
 return send_file(out,as_attachment=True,download_name="FayoumScan_Patients_"+(request.args.get("date") or "All")+".xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
@app.get("/api/price-plans/<int:pid>/export.xlsx")
def exportplan(pid):
 with con() as c:
  p=c.execute("select p.*,t.name payment_type from price_plans p join payment_types t on t.id=p.payment_type_id where p.id=?",(pid,)).fetchone()
  if not p:return jsonify(error="plan not found"),404
  rows=c.execute("select e.unit,e.id exam_id,e.name,i.price,i.coverage_percentage,i.active from exams e left join price_plan_items i on i.exam_id=e.id and i.price_plan_id=? where e.active=1 order by e.unit,e.name",(pid,)).fetchall()
 wb=Workbook();ws=wb.active;ws.title="Price Plan";ws.append(["Unit","Exam ID","Exam Name","Price","Coverage %","Active"])
 for r in rows:ws.append([r["unit"],r["exam_id"],r["name"],r["price"] if r["price"] is not None else "",r["coverage_percentage"] if r["coverage_percentage"] is not None else 0,r["active"] if r["active"] is not None else 1])
 for c in ws[1]:c.fill=PatternFill("solid",fgColor="1B2A6B");c.font=Font(color="FFFFFF",bold=True);c.alignment=Alignment(horizontal="center")
 for col,w in {"A":18,"B":12,"C":36,"D":16,"E":16,"F":12}.items():ws.column_dimensions[col].width=w
 info=wb.create_sheet("Instructions");info.append(["FAYOUM SCAN RIS - Price Plan"]);info.append(["Payment Type",p["payment_type"]]);info.append(["Plan",p["name"]]);info.append(["Version",p["version"]]);info.append(["Valid From",p["valid_from"]]);info.append(["Do not change Exam ID. Enter Price and Coverage %, then import this workbook."])
 out=BytesIO();wb.save(out);out.seek(0);return send_file(out,as_attachment=True,download_name="PricePlan_"+str(pid)+".xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
@app.post("/api/price-plans/<int:pid>/import.xlsx")
def importplan(pid):
 if not admin():return jsonify(error="admin required"),403
 if "file" not in request.files:return jsonify(error="xlsx file required"),400
 try:wb=load_workbook(request.files["file"],data_only=True);ws=wb["Price Plan"]
 except Exception as e:return jsonify(error="invalid xlsx: "+str(e)),400
 parsed=[];errors=[]
 for n,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
  if not any(v is not None and v!="" for v in row):continue
  try:
   exam_id=int(row[1]);price=float(row[3]);coverage=float(row[4] or 0);active=1 if row[5] in (None,"",1,True,"1","Yes","YES") else 0
   if price<0 or not 0<=coverage<=100:raise ValueError("invalid price/coverage")
   parsed.append((pid,exam_id,price,coverage,active))
  except Exception as e:errors.append({"row":n,"error":str(e)})
 if errors:return jsonify(error="validation failed",errors=errors),422
 with con() as c:
  if not c.execute("select 1 from price_plans where id=?",(pid,)).fetchone():return jsonify(error="plan not found"),404
  for x in parsed:c.execute("insert into price_plan_items(price_plan_id,exam_id,price,coverage_percentage,active) values(?,?,?,?,?) on conflict(price_plan_id,exam_id) do update set price=excluded.price,coverage_percentage=excluded.coverage_percentage,active=excluded.active",x)
 return jsonify(ok=True,imported=len(parsed))
@app.get("/api/audit")
def audit():
 if not admin():return jsonify(error="admin required"),403
 with con() as c:return jsonify([dict(x) for x in c.execute("select * from audit_log order by id desc limit 300")])
@app.post("/api/backup")
def backup():
 if not admin():return jsonify(error="admin required"),403
 f="fayoum_scan_"+datetime.now().strftime("%Y%m%d_%H%M%S")+".db";dest=os.path.join(BACKUPS,f)
 with con() as s:d=sqlite3.connect(dest);s.backup(d);d.close()
 return jsonify(ok=True,file=f)
@app.get("/api/backups")
def backups():
 if not admin():return jsonify(error="admin required"),403
 return jsonify(sorted([x for x in os.listdir(BACKUPS) if x.endswith(".db")],reverse=True))
@app.get("/")
def index():return send_from_directory(WEB,"index.html")
@app.get("/<path:p>")
def staticfiles(p):return send_from_directory(WEB,p)

if __name__=="__main__":
 init()
 try:
  from waitress import serve;serve(app,host="0.0.0.0",port=int(os.getenv("PORT","8787")),threads=8)
 except ImportError:app.run(host="0.0.0.0",port=8787,threaded=True)
